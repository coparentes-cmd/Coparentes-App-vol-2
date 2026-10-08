#!/usr/bin/env python3
"""Generate Coparentes security architecture docs (PL + EN) as Word files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT_DIR = Path("/Users/kingastaszewska/Desktop/Coparentes-App-vol-2-main")


def style_doc(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def add_title(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_codeish(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def build_pl() -> Document:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Coparentes — architektura zabezpieczeń")
    add_p(
        doc,
        "Dokument techniczny (programista → programista). Stan zgodny z kodem Flutter "
        "(Coparentes-App-vol-2) oraz backendu Node/Prisma (coparentes-backend). "
        "Opisuje hashowanie haseł, szyfrowanie E2E czatu, szyfrowanie at-rest po stronie serwera "
        "oraz zabezpieczenia systemowe (sesje, transport, limity, uploady, offline).",
    )

    add_h(doc, "1. Model zagrożeń — jedno zdanie")
    add_p(
        doc,
        "Przy kompromitacji samej bazy (bez sekretów env) chronione są: hasła (bcrypt), "
        "tokeny sesji (HMAC z pepperem), treść czatu E2E (klucze klienta) oraz pola "
        "zaszyfrowane KEY_*. Przy kompromitacji serwera z KEY_* operator czyta kalendarz/"
        "finanse/dokumenty/załączniki; treści E2E nadal wymagają hasła użytkownika, "
        "kodu recovery albo odblokowanego urządzenia.",
    )
    add_p(
        doc,
        "Podział odpowiedzialności: tożsamość logowania = serwer (bcrypt). Treść wiadomości "
        "użytkownika w czacie = klient E2E. Kalendarz / finanse / dokumenty / większość PII = "
        "AES-GCM at-rest na serwerze (operator z kluczami env może odszyfrować).",
    )

    add_h(doc, "2. Hashowanie haseł konta")
    add_bullets(
        doc,
        [
            "Biblioteka: bcryptjs, koszt 12 (nie Argon2 — Argon2id jest tylko do owinięcia klucza E2E).",
            "Rejestracja / dołączanie / konto dziecka: bcrypt.hash(password, 12) → User.passwordHash.",
            "Logowanie / re-auth dziecka / zmiana hasła / usuwanie konta: bcrypt.compare.",
            "Reset hasła (w tym reset hasła dziecka przez rodzica): nowy hash bcrypt cost 12; "
            "token resetu w DB to HMAC-SHA256(SESSION_PEPPER|INTEGRITY_SECRET, \"pwreset:\" + token) — "
            "sam wiersz DB bez peppera nie wystarczy.",
            "Polityka: PASSWORD_MIN_LENGTH = 8 (passwordPolicy.js).",
            "Pliki: authService.js, security.js, passwordPolicy.js; workspace child reset → mail do rodziców.",
        ],
    )

    add_h(doc, "3. E2E szyfrowanie czatu (Flutter)")
    add_p(
        doc,
        "Produkcyjna krypto jest w aplikacji (package:cryptography). Backend tylko przechowuje "
        "nieprzezroczyste envelope’y i sealed keys — nie otwiera treści E2E. PoC Node używa "
        "libsodium; aplikacja implementuje równoważny sealed-box design.",
    )
    add_h(doc, "3.1 Tożsamość i owinięcie klucza prywatnego", 2)
    add_bullets(
        doc,
        [
            "Para tożsamości: X25519. PublicKey w DB (User.publicKey). PrivateKey nigdy plaintext na serwerze.",
            "Owinięcie hasłem: Argon2id → AES-256-GCM, envelope v:1 (privateKeyEnvelope). "
            "Domyślnie ~19456 KiB memory, 2 iteracje, parallelism 1, sól 16 B.",
            "Kod odzyskania: ten sam Argon2id+AES-GCM (recoveryKeyEnvelope); tajemnica = kod Crockford "
            "Base32 z 15 losowych bajtów (120 bitów), format XXXX-XXXX-… (bez I/L/O/U).",
            "Lokalny cache odblokowanego klucza: flutter_secure_storage "
            "(coparentes_e2e_unlocked_key_seed_v1). Na web → localStorage — słabszy model zagrożeń.",
        ],
    )
    add_h(doc, "3.2 Klucz wątku i wiadomości", 2)
    add_bullets(
        doc,
        [
            "Każdy wątek: losowy 32-bajtowy klucz AES (thread key), generowany po stronie klienta.",
            "Dostarczenie klucza uczestnikowi: ephemeral X25519 + HKDF-SHA256 "
            "(info = \"coparentes-e2e-threadkey-v1\") + AES-256-GCM — analog crypto_box_seal. "
            "Kopie w tabeli ThreadKey.encryptedKey.",
            "Treść wiadomości: AES-256-GCM pod thread key → {ciphertext, nonce} (base64). "
            "Serwer widzi ciphertext; potem może owinąć całość KEY_MESSAGES do at-rest (nadal bez plaintextu E2E).",
            "Family-sync: rodzic otwiera swój sealed key, re-sealuje pod publicKey dziecka, "
            "POST …/keys/family-sync — dziecko dołącza do istniejącego wątku Rodzina.",
        ],
    )
    add_p(
        doc,
        "Ścieżki Flutter: e2e_crypto_service.dart, e2e_session_service.dart, e2e_key_storage_service.dart.",
    )

    add_h(doc, "4. Szyfrowanie at-rest po stronie serwera (KEY_*)")
    add_bullets(
        doc,
        [
            "Rdzeń: AES-256-GCM, IV 12 B, format enc:v1:{KEY_NAME}:{iv}:{tag}:{ciphertext} (base64url).",
            "Klucze env: base64 → dokładnie 32 bajty; walidacja przy starcie (validateRequiredSecrets). "
            "Bez derivacji z INTEGRITY_SECRET / JWT_SECRET.",
            "KEY_MESSAGES — content wiadomości, attachmentsJson (operator z kluczem widzi JSON załączników "
            "i systemowe wiadomości plaintext; treść użytkownika E2E to nadal ciphertext+nonce).",
            "KEY_FINANCE — title/note/receiptContentBase64 wydatków.",
            "KEY_HEALTH — wydarzenia type=medical oraz dokumenty category=medical.",
            "KEY_GENERAL — pozostały kalendarz, dokumenty nie-medyczne, imię/szkoła dziecka, payload eksportów.",
        ],
    )
    add_p(doc, "Ścieżki: crypto.service.js, finances.js, calendar.js, documents.js, threads.js, exports.js.")

    add_h(doc, "5. Sesje i auth systemowy")
    add_bullets(
        doc,
        [
            "To NIE jest JWT w middleware. Token sesji: crypto.randomBytes(32).hex; w DB "
            "HMAC-SHA256(sessionPepper, token). Pepper: SESSION_PEPPER || INTEGRITY_SECRET.",
            "Dostarczenie: cookie httpOnly coparentes_session (Secure + SameSite=None w prod, path /api) "
            "i/lub Authorization: Bearer. Flutter: flutter_secure_storage (web: SharedPreferences).",
            "Limit: max 5 sesji / użytkownik; TTL SESSION_TTL_DAYS (domyślnie 30).",
            "OTP/2FA (gdy włączone): 6 cyfr, bcrypt cost 12 na kodzie, TTL/próby/cooldown z env; mail Resend.",
            "Trusted device: losowy token, bcrypt(cost 12) w DB; cookie / X-Trusted-Device-Token; "
            "TTL TRUSTED_DEVICE_TTL_DAYS (domyślnie 30).",
        ],
    )

    add_h(doc, "6. Integrity / INTEGRITY_SECRET")
    add_bullets(
        doc,
        [
            "createIntegrityHash: HMAC-SHA256(INTEGRITY_SECRET, JSON.stringify(payload)) → hex "
            "(hash wiadomości, wydatków, manifestów eksportu).",
            "Consent: hash IP = SHA-256(ip:INTEGRITY_SECRET) — bez surowego IP w audycie.",
            "Uwaga implementacyjna: hash jest zapisywany przy write; pełna weryfikacja read-path "
            "nie jest centralnym gate’em — to evidence/tamper-signal, nie zamiennik szyfrowania.",
        ],
    )

    add_h(doc, "7. Transport i hardening HTTP")
    add_bullets(
        doc,
        [
            "HTTPS wymuszane w production / FORCE_HTTPS (redirect + 403 https_required).",
            "Nagłówki: HSTS preload, nosniff, DENY frame, CSP default-src 'self' (własna warstwa + helmet).",
            "CORS: allowlist FRONTEND_URL / CORS_ORIGINS (+ Netlify / domeny Coparentes); credentials.",
            "Rate limit (express-rate-limit, MemoryStore): login 5/5min, register 10/h, forgot/reset, "
            "OTP, user keys/recovery, child password reset — osobne bucket’y.",
        ],
    )

    add_h(doc, "8. Upload plików")
    add_bullets(
        doc,
        [
            "Dokumenty: magic bytes (PDF/JPEG/PNG/WEBP/HEIC/DOC/DOCX/TXT z heurystyką tekstu) — fileSignature.js.",
            "Odrzucenie spoofowanego mimeType / executable.",
            "Paragony finansowe: obecnie bez tej samej bramki magic-byte (świadoma różnica vs dokumenty).",
        ],
    )

    add_h(doc, "9. Offline / lokalnie (Flutter)")
    add_bullets(
        doc,
        [
            "SecureOfflineCodec: AES (package:encrypt, AESMode.sic/CTR) + IV; prefix enc:; klucz 32 B "
            "w flutter_secure_storage (coparentes_offline_aes_key_v1).",
            "Na web kIsWeb → brak codec — cache SharedPreferences plaintext (słabsze).",
            "To osobna warstwa od E2E i od KEY_* serwera.",
        ],
    )

    add_h(doc, "10. Co NIE jest E2E")
    add_bullets(
        doc,
        [
            "Kalendarz, finanse, dokumenty, imię/szkoła dziecka, eksporty — at-rest KEY_*.",
            "Wiadomości systemowe kanałów — plaintext → KEY_MESSAGES.",
            "Załączniki wiadomości w JSON — KEY_MESSAGES, nie sealed pod thread key.",
            "Email logowania, OTP, kod recovery w mailu — plaintext u dostawcy poczty.",
            "Nazwy wyświetlane / role — typowo plaintext lub KEY_GENERAL, nie E2E.",
        ],
    )

    add_h(doc, "11. Mapa warstw (skrót)")
    add_codeish(
        doc,
        "Warstwa A  Hasło konta ........ bcrypt cost 12 (serwer)\n"
        "Warstwa B  Sesja .............. opaque token + HMAC pepper (cookie/Bearer)\n"
        "Warstwa C  Czat E2E ........... X25519 + Argon2id wrap + AES-GCM thread key (klient)\n"
        "Warstwa D  PII w DB ........... AES-256-GCM KEY_* (serwer)\n"
        "Warstwa E  Transport .......... HTTPS + CORS + rate limits + headers\n"
        "Warstwa F  Offline native ..... AES-SIC lokalnie (+ Keystore/Keychain)\n"
        "Warstwa G  Integrity stamp .... HMAC-SHA256 INTEGRITY_SECRET",
    )

    add_h(doc, "12. Pliki źródłowe (orientacyjnie)")
    add_bullets(
        doc,
        [
            "Backend: src/services/crypto.service.js, authService.js, session.js, otp.service.js, "
            "trustedDevice.service.js, utils/security.js, utils/fileSignature.js, middleware/*",
            "Flutter: lib/services/e2e_*.dart, lib/data/local/offline_store.dart, "
            "secure_offline_codec.dart, auth_repository.dart",
        ],
    )

    add_p(
        doc,
        "Dokument wygenerowany automatycznie z opisu architektury zweryfikowanego względem kodu. "
        "Przy zmianie kryptografii zaktualizuj ten plik razem z PR.",
    )
    return doc


def build_en() -> Document:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Coparentes — Security Architecture")
    add_p(
        doc,
        "Engineer-to-engineer technical note. Matches Flutter (Coparentes-App-vol-2) and "
        "Node/Prisma backend (coparentes-backend). Covers password hashing, chat E2E, "
        "server-side at-rest encryption, and system controls (sessions, transport, limits, "
        "uploads, offline cache).",
    )

    add_h(doc, "1. Threat model — one liner")
    add_p(
        doc,
        "DB dump without env secrets: account passwords (bcrypt), session tokens (HMAC pepper), "
        "E2E chat bodies (client keys), and KEY_*-encrypted fields remain protected. "
        "Server operator / stolen KEY_* + DB: calendar, finance, documents, attachments readable; "
        "E2E plaintext still needs the user password, recovery code, or an unlocked device seed.",
    )
    add_p(
        doc,
        "Responsibility split: login identity = server bcrypt. User chat message bodies = client E2E. "
        "Calendar / finance / documents / most PII = server AES-GCM at rest (decryptable with env keys).",
    )

    add_h(doc, "2. Account password hashing")
    add_bullets(
        doc,
        [
            "Library: bcryptjs, cost factor 12 (Argon2id is only used for E2E private-key wrapping).",
            "Register / join / child account: bcrypt.hash(password, 12) → User.passwordHash.",
            "Login / child re-auth / change-password / delete-account: bcrypt.compare.",
            "Password reset (including parent-initiated child reset): new bcrypt cost 12 hash; "
            "reset token stored as HMAC-SHA256(SESSION_PEPPER|INTEGRITY_SECRET, \"pwreset:\" + token).",
            "Policy: PASSWORD_MIN_LENGTH = 8 (passwordPolicy.js).",
            "Code: authService.js, security.js, passwordPolicy.js; child reset emails parents.",
        ],
    )

    add_h(doc, "3. Chat E2E encryption (Flutter)")
    add_p(
        doc,
        "Production crypto lives in the app (package:cryptography). The backend only stores opaque "
        "envelopes and sealed keys — it never opens E2E plaintext. A Node PoC uses libsodium; "
        "the app implements an equivalent sealed-box design.",
    )
    add_h(doc, "3.1 Identity and private-key wrapping", 2)
    add_bullets(
        doc,
        [
            "Identity keypair: X25519. Public key in DB (User.publicKey). Private key never plaintext on server.",
            "Password wrap: Argon2id → AES-256-GCM envelope v:1 (privateKeyEnvelope). "
            "Defaults ~19456 KiB memory, 2 iterations, parallelism 1, 16-byte salt.",
            "Recovery wrap: same Argon2id+AES-GCM (recoveryKeyEnvelope); secret = Crockford Base32 "
            "from 15 random bytes (120 bits), XXXX-XXXX-… (no I/L/O/U).",
            "Local unlocked-key cache: flutter_secure_storage (coparentes_e2e_unlocked_key_seed_v1). "
            "On web this falls back to localStorage — weaker local threat model.",
        ],
    )
    add_h(doc, "3.2 Thread keys and messages", 2)
    add_bullets(
        doc,
        [
            "Per thread: random 32-byte AES thread key, generated on the client.",
            "Key delivery: ephemeral X25519 + HKDF-SHA256 (info = \"coparentes-e2e-threadkey-v1\") "
            "+ AES-256-GCM — crypto_box_seal analogue. Copies in ThreadKey.encryptedKey.",
            "Message body: AES-256-GCM under the thread key → {ciphertext, nonce} (base64). "
            "Server stores ciphertext; may also wrap with KEY_MESSAGES for at-rest (still no E2E plaintext).",
            "Family-sync: parent opens their sealed key, re-seals for the child’s publicKey, "
            "POST …/keys/family-sync so a linked child can join an existing Family thread.",
        ],
    )
    add_p(
        doc,
        "Flutter paths: e2e_crypto_service.dart, e2e_session_service.dart, e2e_key_storage_service.dart.",
    )

    add_h(doc, "4. Server at-rest encryption (KEY_*)")
    add_bullets(
        doc,
        [
            "Core: AES-256-GCM, 12-byte IV, wire format enc:v1:{KEY_NAME}:{iv}:{tag}:{ciphertext} (base64url).",
            "Env keys: base64 → exactly 32 bytes; validated at startup. Not derived from INTEGRITY_SECRET / JWT_SECRET.",
            "KEY_MESSAGES — message content, attachmentsJson (operator with key sees attachment JSON and "
            "plaintext system messages; user E2E bodies remain ciphertext+nonce).",
            "KEY_FINANCE — expense title/note/receiptContentBase64.",
            "KEY_HEALTH — medical calendar events and medical document content.",
            "KEY_GENERAL — other calendar fields, non-medical docs, child name/school, export payloads.",
        ],
    )
    add_p(doc, "Paths: crypto.service.js, finances.js, calendar.js, documents.js, threads.js, exports.js.")

    add_h(doc, "5. Sessions and system auth")
    add_bullets(
        doc,
        [
            "Not JWT in the auth middleware. Session token: crypto.randomBytes(32).hex; DB stores "
            "HMAC-SHA256(sessionPepper, token). Pepper: SESSION_PEPPER || INTEGRITY_SECRET.",
            "Delivery: httpOnly cookie coparentes_session (Secure + SameSite=None in prod, path /api) "
            "and/or Authorization: Bearer. Flutter: flutter_secure_storage (web: SharedPreferences).",
            "Cap: max 5 sessions per user; TTL SESSION_TTL_DAYS (default 30).",
            "OTP/2FA (when enabled): 6-digit code, bcrypt cost 12 on the code; TTL/attempts/cooldown from env; Resend email.",
            "Trusted device: random token, bcrypt(cost 12) in DB; cookie / X-Trusted-Device-Token; "
            "TTL TRUSTED_DEVICE_TTL_DAYS (default 30).",
        ],
    )

    add_h(doc, "6. Integrity / INTEGRITY_SECRET")
    add_bullets(
        doc,
        [
            "createIntegrityHash: HMAC-SHA256(INTEGRITY_SECRET, JSON.stringify(payload)) → hex "
            "(message/expense/export manifest hashes).",
            "Consent IP hashing: SHA-256(ip:INTEGRITY_SECRET).",
            "Implementation note: hashes are written on write; there is no central read-path gate — "
            "tamper evidence, not a substitute for encryption.",
        ],
    )

    add_h(doc, "7. Transport and HTTP hardening")
    add_bullets(
        doc,
        [
            "HTTPS enforced in production / FORCE_HTTPS (redirect + 403 https_required).",
            "Headers: HSTS preload, nosniff, DENY frame, CSP default-src 'self' (custom layer + helmet).",
            "CORS allowlist via FRONTEND_URL / CORS_ORIGINS (plus Netlify / Coparentes domains); credentials.",
            "Rate limits (express-rate-limit, MemoryStore): login, register, forgot/reset, OTP, "
            "user keys/recovery, child password reset — separate buckets.",
        ],
    )

    add_h(doc, "8. File uploads")
    add_bullets(
        doc,
        [
            "Documents: magic-byte sniff (PDF/JPEG/PNG/WEBP/HEIC/DOC/DOCX/TXT with text heuristic) — fileSignature.js.",
            "Rejects spoofed mimeType / executables.",
            "Finance receipts: no equivalent magic-byte gate today (gap vs documents).",
        ],
    )

    add_h(doc, "9. Offline / local (Flutter)")
    add_bullets(
        doc,
        [
            "SecureOfflineCodec: AES via package:encrypt (AESMode.sic/CTR) + IV; enc: prefix; "
            "32-byte key in flutter_secure_storage (coparentes_offline_aes_key_v1).",
            "On web (kIsWeb) codec is disabled — SharedPreferences plaintext (weaker).",
            "Separate from both E2E and server KEY_*.",
        ],
    )

    add_h(doc, "10. What is NOT E2E")
    add_bullets(
        doc,
        [
            "Calendar, finance, documents, child name/school, exports — KEY_* at rest.",
            "System channel messages — plaintext then KEY_MESSAGES.",
            "Message attachments JSON — KEY_MESSAGES, not sealed under the thread key.",
            "Login email, OTP, recovery code email — plaintext at the mail provider.",
            "Display names / roles — typically plaintext or KEY_GENERAL, not E2E.",
        ],
    )

    add_h(doc, "11. Layer map (short)")
    add_codeish(
        doc,
        "Layer A  Account password .... bcrypt cost 12 (server)\n"
        "Layer B  Session ............. opaque token + HMAC pepper (cookie/Bearer)\n"
        "Layer C  Chat E2E ............ X25519 + Argon2id wrap + AES-GCM thread key (client)\n"
        "Layer D  PII in DB ........... AES-256-GCM KEY_* (server)\n"
        "Layer E  Transport ........... HTTPS + CORS + rate limits + headers\n"
        "Layer F  Offline native ...... local AES-SIC (+ Keystore/Keychain)\n"
        "Layer G  Integrity stamp ..... HMAC-SHA256 INTEGRITY_SECRET",
    )

    add_h(doc, "12. Source map")
    add_bullets(
        doc,
        [
            "Backend: src/services/crypto.service.js, authService.js, session.js, otp.service.js, "
            "trustedDevice.service.js, utils/security.js, utils/fileSignature.js, middleware/*",
            "Flutter: lib/services/e2e_*.dart, lib/data/local/offline_store.dart, "
            "secure_offline_codec.dart, auth_repository.dart",
        ],
    )

    add_p(
        doc,
        "Generated from a code-verified architecture summary. Update this document with any crypto PR.",
    )
    return doc


def main() -> None:
    pl_path = OUT_DIR / "Coparentes-Architektura-Zabezpieczen-PL.docx"
    en_path = OUT_DIR / "Coparentes-Security-Architecture-EN.docx"
    build_pl().save(pl_path)
    build_en().save(en_path)
    print(f"Wrote {pl_path}")
    print(f"Wrote {en_path}")


if __name__ == "__main__":
    main()
